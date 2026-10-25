"""Build the two colleague packets (Word), their answer sheets (Excel) and the hidden key.

Reads draw.json (sample_tasks.py), renders/<label>/ (render_tasks.py, packet A) and
renders_audit/<label>/ (render_audit.py, packet B), and the benchmark task list. Writes
  writing/delivery/human_study_2026-10-15/Rewrite_Packet_Block1.docx ... Block5.docx,
                                          Rewrite_Answers_Block1.xlsx ... Block5.xlsx,
                                          Audit_Packet.docx, Audit_Answers.xlsx
  packet_key.json beside this script   (label -> task, arm, scores; the answer key, never sent and not released)
Run: taskset -c 16-23 python build_packets.py
"""
import json, os, re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = ''
OUT = ROOT + 'writing/delivery/human_study_2026-10-15/'
REND_A = os.path.join(HERE, 'renders')
REND_B = os.path.join(HERE, 'renders_audit')
FONT = 'Times New Roman'
BOX = '□'
N_BLOCKS = 5

Q1 = ('Q1', 'Does the AFTER model carry out the request correctly?', ['yes', 'partly', 'no'])
Q2 = ('Q2', 'If not fully: what is wrong?')
Q3 = ('Q3', 'Would you accept this model into the project without further changes?', ['yes', 'no'])


# --------------------------------------------------------------------------- docx helpers
def base_doc():
    d = Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)
    st = d.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    st.paragraph_format.space_after = Pt(3)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(fp, 'PAGE')
    return d


def _field(par, code):
    r = par.add_run()
    for tag, text in (('begin', None), (None, code), ('end', None)):
        if tag:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), tag)
            r._r.append(e)
        else:
            t = OxmlElement('w:instrText')
            t.set(qn('xml:space'), 'preserve')
            t.text = text
            r._r.append(t)


def para(d, text='', size=None, bold=False, italic=False, align=None, after=None, before=None, keep=False):
    p = d.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def bullet(d, text, bold_lead=None):
    p = d.add_paragraph(style='List Bullet')
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)
    return p


def _cell_borders(cell, color='404040', sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:color'), color)
        b.append(e)
    tcPr.append(b)


def _cell_shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill)
    tcPr.append(s)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'nil')
        b.append(e)
    tblPr.append(b)


def boxed(d, label, lines, fill='F2F4F6', size=10.5, bullets=False):
    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    c.width = Cm(17.0)
    _cell_borders(c)
    _cell_shade(c, fill)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label)
    r.bold = True
    r.font.size = Pt(9)
    for ln in ([lines] if isinstance(lines, str) else lines):
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(1.5)
        if bullets:
            p2.paragraph_format.left_indent = Cm(0.35)
            p2.paragraph_format.first_line_indent = Cm(-0.35)
            ln = '•  ' + ln
        r2 = p2.add_run(ln)
        r2.font.size = Pt(size)
    return t


def views(d, rd):
    """BEFORE and AFTER side by side; wall elevations below when present."""
    rows = [('before_plan.png', 'after_plan.png')]
    if os.path.exists(os.path.join(rd, 'before_elev.png')) and os.path.exists(os.path.join(rd, 'after_elev.png')):
        rows.append(('before_elev.png', 'after_elev.png'))
    t = d.add_table(rows=len(rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _no_borders(t)
    for i, (a, b) in enumerate(rows):
        for j, f in enumerate((a, b)):
            c = t.rows[i].cells[j]
            c.width = Cm(8.5)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.add_run().add_picture(os.path.join(rd, f), width=Cm(8.35))
    return len(rows)


def task_heading(d, label):
    p = para(d, label, size=14, bold=True, after=4, keep=True)
    p.paragraph_format.page_break_before = True
    return p


def clean_props(d, title):
    cp = d.core_properties
    cp.author = ''
    cp.last_modified_by = ''
    cp.comments = ''
    cp.title = title


def intro(d, title, items):
    para(d, title, size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    for item in items:
        kind, text = item[0], item[1]
        if kind == 'p':
            para(d, text, after=5)
        elif kind == 'h':
            para(d, text, size=11.5, bold=True, before=6, after=2, keep=True)
        else:
            bullet(d, text, item[2] if len(item) > 2 else None)


LIST = re.compile(r"((?:the (\w+(?: \w+)?) '([^']*)'(?:, | and )){1,}the \w+(?: \w+)? '[^']*')")


def collapse_names(note):
    """'the column 'X', the column 'X' and the column 'X'' -> '3 columns named 'X''."""
    def rep(mo):
        items = re.findall(r"the (\w+(?: \w+)?) '([^']*)'", mo.group(1))
        if len(items) > 1 and len(set(items)) == 1:
            w, n = items[0]
            return f"{len(items)} {w}s named '{n}'"
        return mo.group(1)
    return LIST.sub(rep, note)


# --------------------------------------------------------------------------- texts
def rewrite_intro(block, first, last):
    return [
        ('p', f'Thank you for taking part. This packet holds 40 short tasks, labelled {first} to {last}. Each task '
              'shows one storey of a building model before and after a single change. The changed elements are '
              'outlined in red, and a short note under the drawings says what changed.'),
        ('h', 'What we ask'),
        ('p', 'For each task, write the request you would send to a modeller so that they make exactly this '
              'change. Use your own words, as you would in a normal working day.'),
        ('b', 'Refer to things the way you usually do: storey, room, grid line, a neighbouring wall or column, '
              'element names, compass directions, distances.'),
        ('b', 'The modeller sees the model but not the red outlines, so describe each element so that they can '
              'find it.'),
        ('b', 'Give the values they need to make the change exactly: sizes, distances, and the name of any new '
              'element.'),
        ('b', 'One to three sentences per task are enough. Please do not copy the note word for word.'),
        ('h', 'How to read the drawings'),
        ('b', 'Walls are grey, columns dark grey, doors and windows thin white marks in the walls, floor slabs thin '
              'light grey outlines, and rooms dashed outlines with their names.'),
        ('b', 'Changed elements are outlined in red. A dashed red outline marks an element that was removed. A pale '
              'red patch shows where a moved or reshaped element stood before.'),
        ('b', 'Names in small boxes belong to the element marked by the dot. North points up and the scale bar is '
              'in metres.'),
        ('b', 'For doors, windows and wall heights a face-on view of the wall is added. "Wall start" marks the end '
              'from which positions along the wall are measured.'),
        ('h', 'Practical points'),
        ('b', f'Type your requests into Rewrite_Answers_Block{block}.xlsx, one row per task, and fill in the '
              'short header (initials, years of experience, main modelling tool, date).'),
        ('b', 'Expect about one minute per task, so 40 to 60 minutes in total. You can split the work over '
              'several sittings.'),
        ('b', 'Please work on your own and do not discuss the tasks with anyone until all answers are back.'),
        ('b', 'When you are done, add your initials to the file name (for example '
              f'Rewrite_Answers_Block{block}_AB.xlsx) and send it back.'),
    ]


AUDIT_INTRO = [
    ('p', 'Thank you for taking part. This packet holds 60 short cases, labelled A01 to A60. In each case someone '
          'asked for a change to a building model, and the model was then edited. Please judge the edited model as '
          'you would judge a colleague\'s work before it goes into the project.'),
    ('h', 'What each page shows'),
    ('b', 'The request, word for word.'),
    ('b', 'A plan of the storey BEFORE and AFTER the edit. For doors and windows, and for wall heights, a face-on '
          'view of the wall is added, with sill and head heights.'),
    ('b', 'A short list of the differences between the two model files. It is read from the files, so it also '
          'covers what a plan cannot show (names, types, materials, property values, sizes, storey assignment, '
          'relationships). The list and the red outlines show what changed. They do not say whether the change is '
          'right.'),
    ('h', 'How to read the drawings'),
    ('b', 'Walls are grey, columns dark grey, doors and windows thin white marks in the walls, floor slabs thin '
          'light grey outlines, and rooms dashed outlines with their names.'),
    ('b', 'Elements that differ between the two files are outlined in red. A dashed red outline marks an element '
          'that was removed. A pale red patch shows where a moved or reshaped element stood before.'),
    ('b', 'Names and IDs in small boxes belong to the element marked by the dot; they are the elements the request '
          'names, and the changed elements. North points up and the scale bar is in metres.'),
    ('h', 'Questions for each case'),
    ('b', ' Does the AFTER model carry out the request correctly? (yes / partly / no)', 'Q1'),
    ('b', ' If not fully: what is wrong? (a few words)', 'Q2'),
    ('b', ' Would you accept this model into the project without further changes? (yes / no)', 'Q3'),
    ('h', 'Practical points'),
    ('b', 'Enter your answers in Audit_Answers.xlsx (drop-down lists for Q1 and Q3), or tick the boxes on a '
          'printed copy. Fill in the short header (initials, years of experience, main modelling tool, date).'),
    ('b', 'Expect one to one and a half minutes per case, so 60 to 90 minutes in total. You can split the work over '
          'several sittings.'),
    ('b', 'Please work on your own and do not discuss the cases with anyone until all answers are back.'),
    ('b', 'When you are done, add your initials to the file name (for example Audit_Answers_AB.xlsx) and send it '
          'back.'),
]


def question_lines(d):
    p = para(d, f'{Q1[0]}  {Q1[1]}', size=10, bold=True, after=1, before=5, keep=True)
    para(d, '      ' + '        '.join(f'{BOX} {o}' for o in Q1[2]), size=10.5, after=2, keep=True)
    para(d, f'{Q2[0]}  {Q2[1]}', size=10, bold=True, after=1, before=4, keep=True)
    para(d, '_' * 88, size=10.5, before=8, after=0, keep=True)
    para(d, f'{Q3[0]}  {Q3[1]}', size=10, bold=True, after=1, before=6, keep=True)
    para(d, '      ' + '        '.join(f'{BOX} {o}' for o in Q3[2]), size=10.5, after=0)
    return p


def ruled_lines(d, n=4):
    para(d, 'Your request (one to three sentences):', size=10, bold=True, before=6, after=2, keep=True)
    for _ in range(n):
        para(d, '_' * 88, size=10.5, before=9, after=0)


# --------------------------------------------------------------------------- xlsx
THIN = Side(style='thin', color='999999')
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def header_block(ws, title, extra=()):
    ws['A1'] = title
    ws['A1'].font = Font(name=FONT, bold=True, size=13)
    fields = ['Respondent initials', 'Years of experience in BIM or construction', 'Main modelling tool', 'Date']
    for i, f in enumerate(fields):
        r = 2 + i
        ws.cell(row=r, column=1, value=f).font = Font(name=FONT, bold=True)
        c = ws.cell(row=r, column=2)
        c.fill = PatternFill('solid', fgColor='FFF7D6')
        c.border = BORDER
        c.font = Font(name=FONT)
    r = 2 + len(fields)
    for k, v in extra:
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=2, value=v).font = Font(name=FONT)
        r += 1
    return r + 1      # header row of the table


def table_head(ws, r0, heads, widths):
    for j, h in enumerate(heads):
        c = ws.cell(row=r0, column=1 + j, value=h)
        c.font = Font(name=FONT, bold=True)
        c.fill = PatternFill('solid', fgColor='E3E7EB')
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = BORDER
    for j, w in enumerate(widths):
        ws.column_dimensions[chr(ord('A') + j)].width = w


AUDIT_HEADS = ['Label', 'Q1 Carries out the request correctly?', 'Q2 If not fully: what is wrong?',
               'Q3 Accept without further changes?']
AUDIT_W = (10, 22, 80, 22)


def audit_xlsx(path, labels):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Answers'
    r0 = header_block(ws, 'Check of edited building models: answer sheet')
    table_head(ws, r0, AUDIT_HEADS, AUDIT_W)
    lists = wb.create_sheet('Lists')
    for j, opts in enumerate((Q1[2], Q3[2])):
        for i, o in enumerate(opts):
            lists.cell(row=1 + i, column=1 + j, value=o)
    lists.sheet_state = 'hidden'
    for col, L, n in (('B', 'A', len(Q1[2])), ('D', 'B', len(Q3[2]))):
        dv = DataValidation(type='list', formula1=f'=Lists!${L}$1:${L}${n}', allow_blank=True)
        dv.error = 'Please pick one of the listed answers.'
        dv.errorTitle = 'Answer'
        ws.add_data_validation(dv)
        dv.add(f'{col}{r0 + 1}:{col}{r0 + len(labels)}')
    for i, lab in enumerate(labels):
        r = r0 + 1 + i
        ws.cell(row=r, column=1, value=lab).font = Font(name=FONT, bold=True)
        for j in range(4):
            c = ws.cell(row=r, column=1 + j)
            c.border = BORDER
            if j:
                c.font = Font(name=FONT)
                c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 30
    ws.freeze_panes = f'B{r0 + 1}'
    ex = wb.create_sheet('Example', 1)
    r1 = header_block(ex, 'Example of a filled sheet (for illustration only)')
    for k, v in enumerate(['AB', '8', 'Revit', '2026-10-20']):
        ex.cell(row=2 + k, column=2, value=v)
    table_head(ex, r1, AUDIT_HEADS, AUDIT_W)
    rows = [['A00', 'partly', 'The door is in the right wall, but 0.5 m further from the corner than asked.', 'no'],
            ['A00', 'yes', '', 'yes'],
            ['A00', 'no', 'The wrong column was removed; the one asked for is still there.', 'no']]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = ex.cell(row=r1 + 1 + i, column=1 + j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.properties.creator = ''
    wb.save(path)


REWRITE_HEADS = ['Block', 'Label', 'Your request (one to three sentences)', 'Comment (optional)']
REWRITE_W = (8, 10, 95, 40)


def rewrite_xlsx(path, block, labels):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Answers'
    r0 = header_block(ws, f'Write the change request in your own words: answer sheet, block {block}',
                      extra=[('Block', block)])
    table_head(ws, r0, REWRITE_HEADS, REWRITE_W)
    for i, lab in enumerate(labels):
        r = r0 + 1 + i
        ws.cell(row=r, column=1, value=block).font = Font(name=FONT)
        ws.cell(row=r, column=2, value=lab).font = Font(name=FONT, bold=True)
        for j in range(4):
            c = ws.cell(row=r, column=1 + j)
            c.border = BORDER
            if j >= 2:
                c.font = Font(name=FONT)
                c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 45
    ws.freeze_panes = f'C{r0 + 1}'
    ex = wb.create_sheet('Example', 1)
    r1 = header_block(ex, 'Example of a filled sheet (for illustration only)', extra=[('Block', block)])
    for k, v in enumerate(['AB', '8', 'Revit', '2026-10-20']):
        ex.cell(row=2 + k, column=2, value=v)
    table_head(ex, r1, REWRITE_HEADS, REWRITE_W)
    row = [block, f'R{block}-00', "On Level 2, add a 0.9 x 2.1 m door called 'D-12' in the corridor wall of room 204, "
                                  "about 1.5 m from its west end, at floor level.", '']
    for j, v in enumerate(row):
        c = ex.cell(row=r1 + 1, column=1 + j, value=v)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.properties.creator = ''
    wb.save(path)


# --------------------------------------------------------------------------- main
def main():
    draw = json.load(open(os.path.join(HERE, 'draw.json')))
    blocks = draw['rewrite_blocks']
    cases = sorted(draw['audit_cases'], key=lambda c: c['label'])
    want = {t for b in blocks for t in b} | {c['task_id'] for c in cases}
    recs = {}
    for line in open(ROOT + 'runs_local/bench_v4/tasks.v4c.jsonl'):
        r = json.loads(line)
        if r['task_id'] in want:
            recs[r['task_id']] = r
    os.makedirs(OUT, exist_ok=True)
    key = {'seed': draw['seed'], 'tasks_file': draw['tasks_file'], 'rewrite': {}, 'audit': {}}
    drop = ('batch_pool', 'families', 'wording', 'expected_world_box', 'relation_edges')
    gid = re.compile(r"(?<![0-9A-Za-z_$])[0-9A-Za-z_$]{22}(?![0-9A-Za-z_$])")
    problems = []

    # ---- packet A: one Word file and one sheet per block
    for b, ids in enumerate(blocks, start=1):
        labs = [f'R{b}-{i + 1:02d}' for i in range(len(ids))]
        d = base_doc()
        intro(d, 'Write the change request in your own words', rewrite_intro(b, labs[0], labs[-1]))
        for lab, tid in zip(labs, ids):
            rd = os.path.join(REND_A, lab)
            v = json.load(open(os.path.join(rd, 'view.json')))
            note = collapse_names(v['note'][0].upper() + v['note'][1:]) if v.get('note') else ''
            if not note or gid.search(note):
                problems.append(f'{lab}: change note empty or carries an ID')
            task_heading(d, lab)
            views(d, rd)
            para(d, '', after=2)
            boxed(d, 'What changed', note)
            ruled_lines(d)
            r = recs[tid]
            key['rewrite'][lab] = {
                'block': b, 'task_id': tid, 'operation': r['operation'], 'category': r['category'],
                'version': r['ifc_version'], 'building': r['building_id'], 'edit_kind': r['edit_kind'],
                'family': r['family'], 'instruction': r['instruction'],
                'edit_params': {k: val for k, val in r['edit_params'].items() if k not in drop},
                'change_note': note, 'storey_before': v['storey_before'], 'storey_after': v['storey_after'],
                'red_before': v['before_red'], 'red_after': v['after_red'], 'removed': v['removed'],
                'gold_model': v['gold'], 'gold_sha256_match': v['gold_sha256_match']}
        clean_props(d, 'Write the change request in your own words')
        d.save(OUT + f'Rewrite_Packet_Block{b}.docx')
        rewrite_xlsx(OUT + f'Rewrite_Answers_Block{b}.xlsx', b, labs)

    # ---- packet B
    d = base_doc()
    intro(d, 'Check of edited building models', AUDIT_INTRO)
    for c in cases:
        lab = c['label']
        rd = os.path.join(REND_B, lab)
        v = json.load(open(os.path.join(rd, 'view.json')))
        task_heading(d, lab)
        boxed(d, 'Request', recs[c['task_id']]['instruction'])
        para(d, '', after=1)
        views(d, rd)
        para(d, '', after=1)
        boxed(d, 'Differences between the BEFORE and AFTER model files', v['facts'], size=9.5, bullets=True)
        question_lines(d)
        key['audit'][lab] = {k: c[k] for k in ('task_id', 'arm', 'pool', 'per_task_file', 'operation', 'category',
                                                'version', 'tier', 'edit_kind', 'scores', 'completed',
                                                'near_threshold', 'n_fail_axes', 'committed', 'reply_only',
                                                'stop_reason', 'why')}
        key['audit'][lab]['checker_verdict'] = 'completed' if c['completed'] else 'not completed'
        key['audit'][lab]['edited_file'] = os.path.relpath(c['edited_ifc'], ROOT)
        key['audit'][lab]['facts_shown'] = v['facts']
        key['audit'][lab]['unchanged_file'] = v['unchanged_file']
    clean_props(d, 'Check of edited building models')
    d.save(OUT + 'Audit_Packet.docx')
    audit_xlsx(OUT + 'Audit_Answers.xlsx', [c['label'] for c in cases])
    key['arms'] = {'veribim': 'VeriBIM-9B (grpo_v10b_c10, with the library)',
                   'commercial': 'commercial model claude-sonnet-5-5 with the library note',
                   'base': 'Qwen3.5-9B untrained, with the library note'}
    json.dump(key, open(os.path.join(HERE, 'packet_key.json'), 'w'), indent=1, ensure_ascii=False)
    print('problems:', problems)
    print('written:', sorted(os.listdir(OUT)))


if __name__ == '__main__':
    main()
